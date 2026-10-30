"""Behaviour tests on synthetic scans: python3 -m unittest discover -s tests (from the repo root).

No real card art is used anywhere: tests/synth.py draws abstract cards of true size.
"""

import csv
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import cv2

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
import synth  # noqa: E402

from cropdeck import name as cardname  # noqa: E402


def run(*args):
    """Run `cropdeck <args>` the way a user would, from the repo root."""
    return subprocess.run(
        [sys.executable, "-m", "cropdeck", *map(str, args)],
        cwd=ROOT,
        capture_output=True,
        text=True,
    )


def write_trays(scans: Path, trays):
    scans.mkdir(parents=True, exist_ok=True)
    n = 1
    for fronts, backs in trays:
        cv2.imwrite(str(scans / f"tray_{n:04d}.png"), fronts)
        cv2.imwrite(str(scans / f"tray_{n + 1:04d}.png"), backs)
        n += 2


def close(a, b, tol=40):
    return all(abs(x - y) <= tol for x, y in zip(a, b, strict=True))


class Split(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.scans, self.out = self.tmp / "scans", self.tmp / "sorted"

    def split(self, *extra):
        return run(
            "split", self.scans, "-o", self.out, "--pair", "--expect", 6, "--dpi", synth.DPI, *extra
        )

    def test_six_cards_paired_by_slot_at_true_scale(self):
        f, b, marks = synth.tray(0)
        write_trays(self.scans, [(f, b)])
        r = self.split()
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        fronts = sorted(self.out.glob("card_*_front.png"))
        self.assertEqual(len(fronts), 6)
        for fp in fronts:
            front, back = cv2.imread(str(fp)), cv2.imread(str(fp).replace("_front", "_back"))
            # a front and its back carry the same corner colour only if they are the same card
            self.assertTrue(close(synth.mark_colour(front), synth.mark_colour(back)), fp.name)
            # true scale: 63 x 88 mm plus 1 mm padding each side, at scan dpi, never rescaled
            h, w = front.shape[:2]
            self.assertAlmostEqual(w, round(65 * synth.MM), delta=3)
            self.assertAlmostEqual(h, round(90 * synth.MM), delta=3)
        # six different cards, not one card cut six times
        cols = [synth.mark_colour(cv2.imread(str(p))) for p in fronts]
        self.assertEqual(len({tuple(round(v / 40) for v in c) for c in cols}), 6)

    def test_six_sideways_cards_pair_and_stand_up_at_true_scale(self):
        self.check_layout("2x3-sideways", 6)

    def test_eight_sideways_cards_in_four_rows_pair_and_stand_up_at_true_scale(self):
        self.check_layout("2x4-sideways", 8)  # the most cards A4 takes with room to spare

    def check_layout(self, layout: str, n: int):
        f, b, _ = synth.tray(0, layout=layout)
        write_trays(self.scans, [(f, b)])
        r = run("split", self.scans, "-o", self.out, "--pair", "--expect", n, "--dpi", synth.DPI)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        fronts = sorted(self.out.glob("card_*_front.png"))
        self.assertEqual(len(fronts), n)
        for fp in fronts:
            front, back = cv2.imread(str(fp)), cv2.imread(str(fp).replace("_front", "_back"))
            self.assertTrue(close(synth.mark_colour(front), synth.mark_colour(back)), fp.name)
            h, w = front.shape[:2]
            self.assertAlmostEqual(w, round(65 * synth.MM), delta=3)  # upright: narrow side across
            self.assertAlmostEqual(h, round(90 * synth.MM), delta=3)
        cols = [synth.mark_colour(cv2.imread(str(p))) for p in fronts]
        self.assertEqual(len({tuple(round(v / 40) for v in c) for c in cols}), n)

    def test_resume_cuts_nothing_twice_and_numbers_on(self):
        f, b, _ = synth.tray(0)
        write_trays(self.scans, [(f, b)])
        self.split()
        self.assertEqual(self.split().returncode, 0)
        self.assertEqual(len(list(self.out.glob("card_*_front.png"))), 6)  # second run: nothing new
        f2, b2, _ = synth.tray(1)
        cv2.imwrite(str(self.scans / "tray_0003.png"), f2)
        cv2.imwrite(str(self.scans / "tray_0004.png"), b2)
        self.split()
        self.assertTrue((self.out / "card_0012_front.png").exists())  # numbering continues
        with open(self.out / "cards.csv", newline="", encoding="utf-8") as fh:
            self.assertEqual(len(list(csv.DictReader(fh))), 12)  # manifest appended, not replaced

    def test_a_tray_without_its_backs_is_refused_and_nothing_written(self):
        f, _, _ = synth.tray(0)
        f2, _, _ = synth.tray(1)
        write_trays(self.scans, [(f, f2)])  # the "backs" scan is another tray's fronts
        r = self.split()
        self.assertIn("STOPPED", r.stdout + r.stderr)
        self.assertEqual(r.returncode, 1)
        self.assertEqual(list(self.out.glob("card_*.png")), [])

    def test_a_skipped_tray_is_named_not_reported_as_done(self):
        f, _, _ = synth.tray(0, layout="2x4-sideways")  # 8 fronts ...
        _, b, _ = synth.tray(0, layout="2x3-sideways")  # ... but only 6 backs
        write_trays(self.scans, [(f, b)])
        r = self.split()
        self.assertEqual(r.returncode, 1, r.stdout + r.stderr)
        self.assertIn("Nothing written for 1 tray(s)", r.stdout)
        self.assertIn("tray_0001.png / tray_0002.png", r.stdout)
        self.assertNotIn("nothing new", r.stdout)
        self.assertEqual(list(self.out.glob("card_*.png")), [])

    def test_white_and_dark_backing_both_work(self):
        f, b, _ = synth.tray(2, backing=(25, 25, 25))
        write_trays(self.scans, [(f, b)])
        self.assertEqual(self.split().returncode, 0)
        self.assertEqual(len(list(self.out.glob("card_*_front.png"))), 6)


class Command(unittest.TestCase):
    """The `cropdeck` entry point: run crops and pairs by default, extra steps are opt-in."""

    def test_run_crops_and_pairs_a_folder_and_nothing_more(self):
        home = Path(tempfile.mkdtemp())
        f, b, _ = synth.tray(0, layout="2x4-sideways")  # the default: 8 cards, no --expect
        write_trays(home / "scans", [(f, b)])
        r = run("run", home, "--dpi", synth.DPI)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertNotIn("expected", r.stdout)
        self.assertEqual(len(list((home / "sorted").glob("card_*_front.png"))), 8)
        self.assertEqual(len(list((home / "sorted").glob("card_*_back.png"))), 8)
        self.assertFalse(list(home.glob("*.xlsx")))  # tracker only with --tracker

    def test_status_reports_without_processing(self):
        home = Path(tempfile.mkdtemp())
        f, b, _ = synth.tray(0)
        write_trays(home / "scans", [(f, b)])
        r = run("status", home)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("2 scan(s) waiting", r.stdout)
        self.assertFalse(list((home / "sorted").glob("card_*")))

    def test_cropping_does_not_need_the_optional_packages(self):
        # name needs Tesseract and track needs openpyxl; split must work without either
        code = (
            "import sys; sys.modules['pytesseract'] = None; sys.modules['openpyxl'] = None; "
            "from cropdeck.cli import main; raise SystemExit(main(['split', '--help']))"
        )
        r = subprocess.run([sys.executable, "-c", code], cwd=ROOT, capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)


class Names(unittest.TestCase):
    def test_noisy_ocr_tokens_match_a_name_in_any_language(self):
        index = cardname.load_names()
        name, en, _ = cardname.best_name(["Charrnander", "—"], index, 0.82)  # OCR read 'm' as 'rn'
        self.assertEqual(en, "Charmander")
        name, en, _ = cardname.best_name(["Pummeluff"], index, 0.82)
        self.assertEqual((name, en), ("Pummeluff", "Jigglypuff"))

    def test_short_tokens_are_never_guessed(self):
        index = cardname.load_names()
        self.assertIsNone(cardname.best_name(["Muk", "ab"], index, 0.82)[0])

    @unittest.skipUnless(shutil.which("tesseract"), "needs the Tesseract engine")
    def test_naming_keeps_every_row_and_your_own_columns(self):
        home = Path(tempfile.mkdtemp())
        f, b, _ = synth.tray(0)
        write_trays(home / "scans", [(f, b)])
        self.assertEqual(run("run", home, "--expect", 6, "--dpi", synth.DPI).returncode, 0)
        csv_path = home / "sorted" / "cards.csv"
        with open(csv_path, newline="", encoding="utf-8-sig") as fh:
            rows = list(csv.DictReader(fh))
        rows[0]["my_column"] = "kept"  # a column you added yourself
        rows[0]["notes"] = "mine"  # a base column you filled in
        with open(csv_path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
        r = run("name", home / "sorted", "--timeout", 5)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        with open(csv_path, newline="", encoding="utf-8-sig") as fh:
            after = list(csv.DictReader(fh))
        self.assertEqual(len(after), 6)
        self.assertEqual(after[0]["my_column"], "kept")
        self.assertEqual(after[0]["notes"], "mine")
        self.assertEqual(
            after[0]["scan_front"], rows[0]["scan_front"]
        )  # split's columns survive too

    def test_renamed_files_are_not_matched_again(self):
        self.assertTrue(cardname.FRONT_RE.match("card_0007_front"))
        self.assertIsNone(cardname.FRONT_RE.match("0001_Charmander_9-25_front"))


class Tracker(unittest.TestCase):
    def setUp(self):
        tmp = Path(tempfile.mkdtemp())
        self.sorted, self.xlsx = tmp / "sorted", tmp / "t.xlsx"
        self.sorted.mkdir()
        self.cards(
            [
                {"card": "0001", "name": "Pikachu", "name_en": "Pikachu", "number": "25/102"},
                {"card": "0002", "name": "Pikachv", "name_en": "", "number": ""},  # OCR slip
            ]
        )

    def cards(self, rows):
        for r in rows:
            r.update(
                confidence="1.00", front=f"{r['card']}_front.png", back=f"{r['card']}_back.png"
            )
        with open(self.sorted / "cards.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)

    def track(self, *extra):
        r = run("track", self.sorted, "-o", self.xlsx, *extra)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        return r

    def book(self):
        from openpyxl import load_workbook

        return load_workbook(self.xlsx)

    def test_your_edits_survive_a_rerun(self):
        self.track()
        wb = self.book()
        ws = wb["Inventory"]
        head = {c.value: c.column for c in ws[1]}
        ws.cell(row=2, column=head["Buyer / note"], value="my note")  # a column that is yours
        ws.cell(row=2, column=head["Status"], value="Sold")  # one the tracker pre-fills
        wb.save(self.xlsx)
        self.track()
        ws = self.book()["Inventory"]
        self.assertEqual(ws.cell(row=2, column=head["Buyer / note"]).value, "my note")
        self.assertEqual(ws.cell(row=2, column=head["Status"]).value, "Sold")

    def test_a_box_size_once_set_is_kept(self):
        self.track("--box-size", "50")
        self.track()  # no --box-size: the workbook's own setting stays
        self.assertEqual(self.book()["Lists"]["B2"].value, 50)

    def test_duplicates_come_from_your_corrections_in_the_inventory(self):
        self.track()
        wb = self.book()
        ws = wb["Inventory"]
        head = {c.value: c.column for c in ws[1]}
        ws.cell(row=3, column=head["Name"], value="Pikachu")  # you fix the OCR slip ...
        ws.cell(row=3, column=head["Number"], value="25/102")  # ... and it is a second copy
        wb.save(self.xlsx)
        r = self.track()
        self.assertIn("1 card(s) you own more than once", r.stdout)
        values = [c.value for row in self.book()["Duplicates"].iter_rows() for c in row]
        self.assertIn("0001, 0002", values)

    def test_a_new_workbook_is_english_in_euro_by_default(self):
        self.track()
        wb = self.book()
        self.assertEqual(wb.sheetnames, ["Overview", "Inventory", "Duplicates", "Lists"])
        self.assertIn("Price EUR", [c.value for c in wb["Inventory"][1]])
        self.assertIn("Cardmarket", [c.value for c in wb["Lists"]["H"]])

    def test_a_german_workbook_stays_german_in_its_currency(self):
        self.track("--lang", "de", "--currency", "CHF")
        wb = self.book()
        self.assertEqual(wb.sheetnames, ["Übersicht", "Inventar", "Dubletten", "Listen"])
        self.assertIn("Preis CHF", [c.value for c in wb["Inventar"][1]])
        self.assertIn("Ricardo", [c.value for c in wb["Listen"]["H"]])
        # an older workbook has no stored settings: language and currency are read from it
        wb["Listen"]["B3"], wb["Listen"]["B4"] = None, None
        wb.save(self.xlsx)
        r = self.track("--lang", "en")
        self.assertIn("stays so - en ignored", r.stdout)
        wb = self.book()
        self.assertEqual(wb.sheetnames, ["Übersicht", "Inventar", "Dubletten", "Listen"])
        self.assertIn("Preis CHF", [c.value for c in wb["Inventar"][1]])

    def test_every_language_has_exactly_the_texts_english_has(self):
        from cropdeck import texts

        def shape(d):
            return {k: shape(v) if isinstance(v, dict) else type(v) for k, v in d.items()}

        for code in texts.languages():
            self.assertEqual(shape(texts.language(code)), shape(texts.language("en")), code)
        for code, cur in texts.currencies().items():
            self.assertTrue(cur["platforms"], code)


if __name__ == "__main__":
    unittest.main()
